"""Интеграционные тесты FastAPI-сервиса (требуют PostgreSQL)."""

from __future__ import annotations

import asyncio
import io
import json
import os
import shutil
import time
from collections.abc import AsyncIterator, Iterator
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, cast

import pytest
from fastapi.testclient import TestClient
from sqlalchemy.ext.asyncio import create_async_engine

from zakupki_parser.api.app import create_app
from zakupki_parser.api.app.routes import procurements as procurements_routes
from zakupki_parser.auth import ROLE_ADMIN, ROLE_ANALYST, ROLE_DEVOPS, ROLE_USER, create_token
from zakupki_parser.config.models import DbConfig
from zakupki_parser.parser.by_url import ProcurementUrlError
from zakupki_parser.storage.db import Base, Database
from zakupki_parser.storage.repository import ProcurementRepository

TEST_DSN = os.environ.get("ZAKUPKI_TEST_DSN", "")
AUTH_SECRET = "test-secret"
# Служебные эндпоинты конвейера (POST /score, /customers/{id}/rating) закрыты
# внутренним токеном (X-Internal-Token) и не принимают пользовательский bearer.
INTERNAL_HEADERS = {"X-Internal-Token": "internal-123"}
# Компетенции — всегда канонический JSON схемы Profile (BR-07): raw-строка не
# проходит валидацию при сохранении профиля в сиде.
COMP_JSON = json.dumps(
    {
        "positioning": "Тестовые компетенции",
        "breadth": "broad",
        "competencies": [{"area": "Аудит", "description": "обследование"}],
        "exclusions": [],
    },
    ensure_ascii=False,
    separators=(",", ":"),
)


async def _add_region_procurement() -> None:
    """Сохраняет закупку с регионом (для проверки region в API/фильтре)."""
    db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
    await db.connect()
    try:
        repo = ProcurementRepository(db)
        await repo.upsert(
            {
                "number": "RG-API-1",
                "platform_id": "zakupki_mos",
                "subject": "Закупка Москвы",
                "region": "г. Москва",
            }
        )
        pid = await repo.find_id("RG-API-1", "zakupki_mos")
        assert pid is not None
    finally:
        await db.dispose()
    await _match_active_profile(pid)


async def _add_region_score_procurement(number: str, region: str) -> int:
    """Закупка с известным регионом на площадке вне конфига (досборка деталей no-op)."""
    db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
    await db.connect()
    try:
        repo = ProcurementRepository(db)
        await repo.upsert(
            {
                "number": number,
                "platform_id": "no-region-platform",
                "subject": "Закупка",
                "region": region,
            }
        )
        pid = await repo.find_id(number, "no-region-platform")
        assert pid is not None
        return pid
    finally:
        await db.dispose()


async def _has_evaluation(procurement_id: int, profile_id: int) -> bool:
    """Существует ли per-profile оценка пары (закупка, профиль)."""
    db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
    await db.connect()
    try:
        repo = ProcurementRepository(db)
        return await repo.get_score(procurement_id, profile_id) is not None
    finally:
        await db.dispose()


async def _match_active_profile(procurement_id: int) -> None:
    """Отмечает закупку как отобранную дефолтным профилем admin (BR-07).

    ``GET /api/procurements`` показывает только закупки, отобранные активным
    профилем (есть строка в ``procurement_evaluations`` — см.
    ``list_procurements``). Закупки, вставленные в тестах напрямую через
    ``repo.upsert``/``ProcurementRepository`` в обход обычного матчинга
    профилем, нужно явно «отобрать», иначе они не попадут в выдачу.
    """
    db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
    await db.connect()
    try:
        repo = ProcurementRepository(db)
        user = await repo.first_user()
        assert user is not None
        profile = await repo.get_active_profile(user.id)
        assert profile is not None and profile.id is not None
        await repo.record_matched_keywords(procurement_id, profile.id, ["тест"])
    finally:
        await db.dispose()


@pytest.fixture(scope="module")
def analyst_headers() -> dict[str, str]:
    """Bearer-токен пользователя с ролью analyst для конфиг/промпт-эндпоинтов."""

    async def _mk() -> int:
        engine = create_async_engine(TEST_DSN)
        async with engine.begin() as conn:
            await conn.run_sync(Base.metadata.create_all)
        await engine.dispose()
        db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
        await db.connect()
        try:
            repo = ProcurementRepository(db)
            user = await repo.create_user("analyst", "test-hash", [ROLE_ANALYST, ROLE_USER])
            return user.id
        finally:
            await db.dispose()

    user_id = asyncio.run(_mk())
    token = create_token(user_id, [ROLE_ANALYST, ROLE_USER], AUTH_SECRET, 3600)
    return {"Authorization": f"Bearer {token}"}


pytestmark = pytest.mark.skipif(not TEST_DSN, reason="ZAKUPKI_TEST_DSN не задан")


@pytest.fixture(scope="module")
def api_client(tmp_path_factory: pytest.TempPathFactory) -> Iterator[tuple[TestClient, Path]]:
    async def _setup() -> int:
        engine = create_async_engine(TEST_DSN)
        async with engine.begin() as conn:
            await conn.run_sync(Base.metadata.drop_all)
            await conn.run_sync(Base.metadata.create_all)
        await engine.dispose()
        # Сид пользователя и его default-профиля.
        db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
        await db.connect()
        try:
            repo = ProcurementRepository(db)
            user = await repo.first_user()
            if user is None:
                user = await repo.create_user(
                    "admin", "test-hash", [ROLE_ADMIN, ROLE_USER, ROLE_DEVOPS]
                )
            # Как начальный администратор: активный аккаунт со всеми платными опциями.
            await repo.ensure_default_account(user.id, paid_default=True)
            await repo.upsert_profile(
                {
                    "name": "default",
                    "enabled": True,
                    "is_active": True,
                    "competencies": COMP_JSON,
                    "keywords": [],
                    "exclusion_words": [],
                },
                user.id,
            )
            return user.id
        finally:
            await db.dispose()

    user_id = asyncio.run(_setup())
    docs = tmp_path_factory.mktemp("docs")

    os.environ["ZAKUPKI_DB_DSN"] = TEST_DSN
    # Авторизация всегда включена: задаём секрет и внутренний токен (обязательны).
    os.environ["ZAKUPKI_AUTH_SECRET"] = AUTH_SECRET
    os.environ["ZAKUPKI_INTERNAL_TOKEN"] = "internal-123"
    app = create_app()
    with TestClient(app) as client:
        token = create_token(user_id, [ROLE_ADMIN, ROLE_USER, ROLE_DEVOPS], AUTH_SECRET, 3600)
        client.headers["Authorization"] = f"Bearer {token}"
        yield client, docs
    os.environ.pop("ZAKUPKI_DB_DSN", None)
    os.environ.pop("ZAKUPKI_AUTH_SECRET", None)
    os.environ.pop("ZAKUPKI_INTERNAL_TOKEN", None)


@pytest.fixture(scope="module")
async def inserted_id(api_client: tuple[TestClient, Path]) -> AsyncIterator[int]:
    client, _ = api_client
    db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
    await db.connect()
    repo = ProcurementRepository(db)
    await repo.upsert(
        {
            "number": "API-1",
            "platform_id": "zakupki_mos",
            "subject": "Тест API",
            "customer": "Заказчик ООО",
            "okpd2_codes": "62.01",
        }
    )
    rows, _ = await repo.list_procurements(number="API-1")
    pid = rows[0].id
    await db.dispose()
    await _match_active_profile(pid)
    yield pid


@pytest.mark.slow  # первый тест модуля — оплачивает setup module-scoped api_client
def test_health(api_client: tuple[TestClient, Path]) -> None:
    client, _ = api_client
    resp = client.get("/health")
    assert resp.status_code == 200
    body = resp.json()
    assert body["status"] == "ok"
    assert body["db"] is True


@pytest.mark.slow  # test_health (обычный «первый») уже slow и деселектится
# под -m "not slow" — setup module-scoped api_client платит ЭТОТ.
def test_coverage(api_client: tuple[TestClient, Path]) -> None:
    """GET /api/coverage — статика (конфиг) + динамика (БД) по площадкам."""
    client, _ = api_client
    resp = client.get("/api/coverage")
    assert resp.status_code == 200
    body = resp.json()
    platforms = body["platforms"]
    assert isinstance(platforms, list) and platforms
    mos = next(p for p in platforms if p["platform_id"] == "zakupki_mos")
    assert 0.0 <= mos["coverage_score"] <= 1.0
    assert isinstance(mos["static"], list) and mos["static"]
    assert all("key" in f and "tier" in f and "status" in f for f in mos["static"])
    # Динамика: пустая БД -> runtime None либо корректная структура.
    assert "runtime" in mos


def test_parser_status_initial(api_client: tuple[TestClient, Path]) -> None:
    client, _ = api_client
    body = client.get("/api/parser/status").json()
    assert body["running"] is False
    assert body["error"] is None


def test_parser_stop_when_idle(api_client: tuple[TestClient, Path]) -> None:
    client, _ = api_client
    resp = client.post("/api/parser/stop")
    assert resp.status_code == 200
    assert resp.json()["status"] == "idle"


def test_restart_process_requires_idle_monitoring(api_client: tuple[TestClient, Path]) -> None:
    """Полный рестарт процесса недоступен, пока идёт обход площадок (409)."""
    client, _ = api_client

    class _FakeRunningTask:
        def done(self) -> bool:
            return False

    state = cast(Any, client.app).state.parser
    state.parser_task = _FakeRunningTask()
    try:
        resp = client.post("/api/parser/restart-process")
        assert resp.status_code == 409
    finally:
        state.parser_task = None


def test_restart_process_schedules_execv_when_idle(
    monkeypatch: pytest.MonkeyPatch, api_client: tuple[TestClient, Path]
) -> None:
    """При простое запрос принимается и (с небольшой задержкой) вызывает os.execv."""
    calls: list[list[str]] = []

    def fake_execv(path: str, argv: list[str]) -> None:
        calls.append(argv)

    monkeypatch.setattr("zakupki_parser.api.app.routes.admin.os.execv", fake_execv)
    client, _ = api_client
    resp = client.post("/api/parser/restart-process")
    assert resp.status_code == 200
    assert resp.json()["status"] == "restarting_process"
    time.sleep(0.6)
    assert len(calls) == 1


def test_db_clear_when_idle(api_client: tuple[TestClient, Path]) -> None:
    client, _ = api_client
    resp = client.post("/api/db/clear")
    assert resp.status_code == 200
    body = resp.json()
    assert body["status"] == "cleared"
    assert client.get("/api/procurements").json()["total"] == 0


def test_db_clear_inactive(api_client: tuple[TestClient, Path]) -> None:
    """POST /api/db/clear-inactive удаляет только неактивные закупки."""
    client, _ = api_client

    async def _seed() -> tuple[int, int]:
        db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
        await db.connect()
        try:
            repo = ProcurementRepository(db)
            await repo.upsert(
                {
                    "number": "CIN-1",
                    "platform_id": "zakupki_mos",
                    "subject": "Неактивна",
                    "is_active": False,
                }
            )
            await repo.upsert(
                {
                    "number": "CIN-2",
                    "platform_id": "zakupki_mos",
                    "subject": "Активна",
                }
            )
            rows, _ = await repo.list_procurements(number="CIN-")
            ids = {p.number: p.id for p in rows}
            return ids["CIN-1"], ids["CIN-2"]
        finally:
            await db.dispose()

    inactive_id, active_id = asyncio.run(_seed())

    resp = client.post("/api/db/clear-inactive")
    assert resp.status_code == 200
    assert resp.json()["deleted"] >= 1

    assert client.get(f"/api/procurements/{inactive_id}").status_code == 404
    assert client.get(f"/api/procurements/{active_id}").status_code == 200


def test_db_clear_irrelevant(api_client: tuple[TestClient, Path]) -> None:
    """POST /api/db/clear-irrelevant удаляет закупки с fit_score < порога (по умолчанию 0.4)."""
    client, _ = api_client

    async def _seed() -> tuple[int, int, int]:
        db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
        await db.connect()
        try:
            repo = ProcurementRepository(db)
            await repo.upsert(
                {
                    "number": "CIR-1",
                    "platform_id": "zakupki_mos",
                    "subject": "Релевантная",
                }
            )
            await repo.upsert(
                {
                    "number": "CIR-2",
                    "platform_id": "zakupki_mos",
                    "subject": "Нерелевантная",
                }
            )
            # Отсечка по векторной близости (ADR-8): fit_score=0 — нерелевантна.
            await repo.upsert(
                {
                    "number": "CIR-3",
                    "platform_id": "zakupki_mos",
                    "subject": "Векторная отсечка",
                }
            )
            rows, _ = await repo.list_procurements(number="CIR-")
            ids = {p.number: p.id for p in rows}
            # Оценки per-profile (BR-07): результат внешнего скоринга приходит
            # через POST /score и пишется в procurement_evaluations активного профиля.
            user = await repo.first_user()
            assert user is not None
            profile = await repo.get_active_profile(user.id)
            assert profile is not None
            await repo.upsert_score(ids["CIR-1"], profile.id, fit_score=0.8, score_method="fit")
            await repo.upsert_score(ids["CIR-2"], profile.id, fit_score=0.2, score_method="fit")
            await repo.upsert_score(ids["CIR-3"], profile.id, fit_score=0.0, score_method="sim")
            return ids["CIR-1"], ids["CIR-2"], ids["CIR-3"]
        finally:
            await db.dispose()

    relevant_id, irrelevant_id, sim_id = asyncio.run(_seed())

    resp = client.post("/api/db/clear-irrelevant", json={"min_fit_score": 0.4})
    assert resp.status_code == 200
    assert resp.json()["deleted"] >= 1

    assert client.get(f"/api/procurements/{relevant_id}").status_code == 200
    assert client.get(f"/api/procurements/{irrelevant_id}").status_code == 404
    assert client.get(f"/api/procurements/{sim_id}").status_code == 404


def test_websocket_receives_broadcast(api_client: tuple[TestClient, Path]) -> None:
    client, _ = api_client
    # Авторизация всегда включена: /ws требует токен query-параметром ?token=
    # (браузер не может задать заголовок WebSocket-запроса).
    token = client.headers["Authorization"].removeprefix("Bearer ")
    with client.websocket_connect(f"/ws?token={token}") as ws:
        # Запрос, меняющий БД, шлёт клиенту "data-changed".
        r = client.post("/api/db/clear")
        assert r.status_code == 200
        assert ws.receive_text() == "data-changed"


def test_index_page_served(api_client: tuple[TestClient, Path]) -> None:
    client, _ = api_client
    resp = client.get("/")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/html")
    assert "Парсер закупок" in resp.text
    assert 'src="/static/js/main.js"' in resp.text
    # token_storage (config_ops.yaml -> auth) подмешивается ДО модульных
    # <script> — api.js читает window.__TOKEN_STORAGE__ синхронно при загрузке.
    assert 'window.__TOKEN_STORAGE__="local"' in resp.text
    assert resp.text.index("__TOKEN_STORAGE__") < resp.text.index('src="/static/js/main.js"')


def test_list_and_get(api_client: tuple[TestClient, Path], inserted_id: int) -> None:
    client, _ = api_client
    resp = client.get("/api/procurements", params={"number": "API-1"})
    assert resp.status_code == 200
    body = resp.json()
    assert body["total"] >= 1
    assert any(item["id"] == inserted_id for item in body["items"])

    detail = client.get(f"/api/procurements/{inserted_id}")
    assert detail.status_code == 200
    assert detail.json()["number"] == "API-1"
    assert detail.json()["is_active"] is True


def test_list_filter_active(api_client: tuple[TestClient, Path], inserted_id: int) -> None:
    client, _ = api_client
    active = client.get("/api/procurements", params={"active": True}).json()
    inactive = client.get("/api/procurements", params={"active": False}).json()
    assert any(item["id"] == inserted_id for item in active["items"])
    assert all(item["is_active"] is True for item in active["items"])
    assert all(item["is_active"] is False for item in inactive["items"])


def test_missing_procurement_404(api_client: tuple[TestClient, Path]) -> None:
    client, _ = api_client
    assert client.get("/api/procurements/999999").status_code == 404


def test_procurement_tz_text(
    api_client: tuple[TestClient, Path], monkeypatch: pytest.MonkeyPatch
) -> None:
    """GET /api/procurements/{id}/tz возвращает текст ТЗ (в т.ч. из архива).

    Кэш задействован через настоящий resolve_tz_content_cached: подменяется только
    извлечение (extract_text) и поиск файла (find_tz_reference), чтобы не ходить
    в сеть. Повторный запрос не переизвлекает текст (счётчик вызовов не растёт).
    """
    from scoring_common.tz import clear_tz_text_cache
    from scoring_common.tz.files import FileRef

    client, _ = api_client
    clear_tz_text_cache()
    try:

        async def _seed() -> int:
            db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
            await db.connect()
            try:
                repo = ProcurementRepository(db)
                assert await repo.upsert(
                    {
                        "number": "TZ-1",
                        "platform_id": "zakupki_mos",
                        "subject": "Закупка с ТЗ в архиве",
                        "files_json": [
                            {"name": "приложение.zip", "url": "http://x/a.zip"},
                            {"name": "смета.xlsx", "url": "http://x/smeta.xlsx"},
                        ],
                    }
                )
                rows, _ = await repo.list_procurements(number="TZ-1")
                return rows[0].id
            finally:
                await db.dispose()

        tz_id = asyncio.run(_seed())

        extract_calls: list[tuple[str, str]] = []

        def fake_find(
            record: dict[str, Any], timeout: float = 30.0, verify_ssl: bool = True
        ) -> FileRef | None:
            files = record.get("files_json") or []
            assert any(f.get("name") == "приложение.zip" for f in files)
            return FileRef("ТЗ.docx", "http://x/a.zip#doc/ТЗ.docx")

        def fake_extract(
            ref: FileRef, timeout: float = 30.0, verify_ssl: bool = True
        ) -> str | None:
            extract_calls.append((ref.url, ref.name))
            # Текст должен содержать требование к Исполнителю, иначе resolve_tz_content
            # уйдёт в фолбэк на документ «Описание» (не подменённый) — см. _has_executor_duties.
            return "# Раздел 1\nИсполнитель обязан поставить товар."

        monkeypatch.setattr("scoring_common.tz.find_tz_reference", fake_find)
        monkeypatch.setattr("scoring_common.tz.extract_text", fake_extract)

        body = client.get(f"/api/procurements/{tz_id}/tz").json()
        assert body["found"] is True
        assert body["file_name"] == "ТЗ.docx"
        assert body["from_archive"] is True
        assert "Раздел 1" in body["text"]
        assert extract_calls == [("http://x/a.zip#doc/ТЗ.docx", "ТЗ.docx")]

        # Повторный запрос отдаёт тот же результат без повторного извлечения (кэш).
        again = client.get(f"/api/procurements/{tz_id}/tz").json()
        assert again["text"] == body["text"]
        assert extract_calls == [("http://x/a.zip#doc/ТЗ.docx", "ТЗ.docx")]
    finally:
        clear_tz_text_cache()


def test_procurement_tz_not_found(api_client: tuple[TestClient, Path]) -> None:
    """Без файлов ТЗ эндпоинт отдаёт found=False (без обращения к сети)."""
    client, _ = api_client

    async def _seed() -> int:
        db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
        await db.connect()
        try:
            repo = ProcurementRepository(db)
            assert await repo.upsert(
                {
                    "number": "TZ-NONE",
                    "platform_id": "zakupki_mos",
                    "subject": "Без ТЗ",
                    "files_json": [{"name": "смета.xlsx", "url": "http://x/smeta.xlsx"}],
                }
            )
            rows, _ = await repo.list_procurements(number="TZ-NONE")
            return rows[0].id
        finally:
            await db.dispose()

    tz_id = asyncio.run(_seed())
    body = client.get(f"/api/procurements/{tz_id}/tz").json()
    assert body["found"] is False
    assert body["text"] is None


def test_procurement_requirements_post_internal(api_client: tuple[TestClient, Path]) -> None:
    """POST /requirements (внутренний, воркер анализа) сохраняет структуру.

    Извлечение (``scoring_common.requirements.extract_requirements``, вызывается
    воркером, не этим API) тестируется отдельно — ``src/scoring_common/tests/
    test_requirements.py``. Здесь — только персист + видимость в detail-ответе
    (единый отчёт читает ``requirements_json`` оттуда, GET .../requirements
    убран вместе с кнопкой «Требования к участнику», FR-13).
    """
    client, _ = api_client

    async def _seed() -> int:
        db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
        await db.connect()
        try:
            repo = ProcurementRepository(db)
            assert await repo.upsert(
                {
                    "number": "REQ-POST",
                    "platform_id": "zakupki_mos",
                    "subject": "Внутренний персист",
                    "files_json": [],
                }
            )
            rows, _ = await repo.list_procurements(number="REQ-POST")
            return rows[0].id
        finally:
            await db.dispose()

    req_id = asyncio.run(_seed())
    structure = {
        "licenses": [{"text": "Требуется лицензия МЧС.", "data": {"required": True}}],
        "other": [{"text": "Состав заявки.", "data": {"type": "состав заявки"}}],
    }
    resp = client.post(
        f"/api/procurements/{req_id}/requirements",
        json={"structure": structure},
        headers=INTERNAL_HEADERS,
    )
    assert resp.status_code == 200
    assert resp.json()["requirements"] == structure

    detail = client.get(f"/api/procurements/{req_id}").json()
    assert detail["requirements_json"] == structure


def test_procurement_requirements_post_empty_object(api_client: tuple[TestClient, Path]) -> None:
    """Пустая структура ({}) сохраняется как есть (не NULL) — виден в detail."""
    client, _ = api_client

    async def _seed() -> int:
        db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
        await db.connect()
        try:
            repo = ProcurementRepository(db)
            assert await repo.upsert(
                {
                    "number": "REQ-NONE",
                    "platform_id": "zakupki_mos",
                    "subject": "Без требований",
                    "files_json": [],
                }
            )
            rows, _ = await repo.list_procurements(number="REQ-NONE")
            return rows[0].id
        finally:
            await db.dispose()

    req_id = asyncio.run(_seed())
    resp = client.post(
        f"/api/procurements/{req_id}/requirements",
        json={"structure": {}},
        headers=INTERNAL_HEADERS,
    )
    assert resp.status_code == 200
    assert resp.json()["requirements"] == {}

    detail = client.get(f"/api/procurements/{req_id}").json()
    assert detail["requirements_json"] == {}


def test_export_procurement_xlsx_highlights_blocking_rows(
    api_client: tuple[TestClient, Path],
) -> None:
    """Excel-экспорт карточки (FR-13.4): вердикт/требования/geo — красным, если блокируют."""
    import openpyxl

    client, _ = api_client

    async def _seed() -> int:
        db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
        await db.connect()
        try:
            repo = ProcurementRepository(db)
            assert await repo.upsert(
                {
                    "number": "XLSX-1",
                    "platform_id": "zakupki_mos",
                    "subject": "Экспорт с блокировкой",
                }
            )
            rows, _ = await repo.list_procurements(number="XLSX-1")
            pid = rows[0].id
            await repo.save_requirements(
                pid, {"licenses": [{"text": "Требуется лицензия МЧС", "data": None}]}
            )
            user = await repo.first_user()
            assert user is not None
            profile = await repo.get_active_profile(user.id)
            assert profile is not None
            await repo.update_rag_report(
                pid,
                profile.id,
                {
                    "verdict": {
                        "accepted": False,
                        "blocking_reasons": [{"source": "licenses", "label": "Лицензии"}],
                    },
                    "requirements_verdict": {
                        "licenses": {"blocking": True, "negated": False, "count": 1}
                    },
                    "geo": {
                        "too_far": True,
                        "distance_km": 120.0,
                        "max_distance_km": 50.0,
                        "region": "Московская область",
                    },
                },
            )
            return pid
        finally:
            await db.dispose()

    pid = asyncio.run(_seed())
    resp = client.get(f"/api/procurements/{pid}/export.xlsx")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert ws is not None
    rows_by_label = {row[0]: row[1] for row in ws.iter_rows(min_row=2, values_only=True)}
    assert rows_by_label["Вердикт приемлемости"] == "Отклонена: Лицензии"
    assert rows_by_label["Лицензии"] == "Требуется лицензия МЧС"
    assert "120.0" in str(rows_by_label["Расстояние до центра региона"])

    # Блокирующие строки — красным (проверяем цвет шрифта конкретных ячеек).
    label_cells = {cell.value: cell for row in ws.iter_rows(min_row=2) for cell in [row[0]]}
    assert label_cells["Вердикт приемлемости"].font.color.rgb == "FFDC2626"
    assert label_cells["Лицензии"].font.color.rgb == "FFDC2626"
    assert label_cells["Расстояние до центра региона"].font.color.rgb == "FFDC2626"
    # Неблокирующая строка (обычные данные закупки) — не подсвечена.
    assert label_cells["Номер"].font.bold is not True


def test_export_procurement_xlsx_license_summary(api_client: tuple[TestClient, Path]) -> None:
    """Excel-экспорт: лицензии — компактной сводкой (вид + наличие у поставщика),
    а не сырым текстом требования (``rag_report.requirements_status.licenses``)."""
    import openpyxl

    client, _ = api_client

    async def _seed() -> int:
        db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
        await db.connect()
        try:
            repo = ProcurementRepository(db)
            assert await repo.upsert(
                {"number": "XLSX-LIC", "platform_id": "zakupki_mos", "subject": "Сводка лицензий"}
            )
            rows, _ = await repo.list_procurements(number="XLSX-LIC")
            pid = rows[0].id
            await repo.save_requirements(
                pid, {"licenses": [{"text": "Требуется лицензия МЧС", "data": None}]}
            )
            user = await repo.first_user()
            assert user is not None
            profile = await repo.get_active_profile(user.id)
            assert profile is not None
            await repo.update_rag_report(
                pid,
                profile.id,
                {
                    "requirements_status": {
                        "licenses": {
                            "found": True,
                            "required": True,
                            "negated": False,
                            "items": [
                                {
                                    "label": "Лицензия МЧС (пожарная безопасность)",
                                    "kind": "mchs",
                                    "available": False,
                                }
                            ],
                        }
                    },
                    "requirements_verdict": {
                        "licenses": {"blocking": False, "negated": False, "count": 1}
                    },
                },
            )
            return pid
        finally:
            await db.dispose()

    pid = asyncio.run(_seed())
    resp = client.get(f"/api/procurements/{pid}/export.xlsx")
    assert resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert ws is not None
    rows_by_label = {row[0]: row[1] for row in ws.iter_rows(min_row=2, values_only=True)}
    assert rows_by_label["Лицензии"] == "Лицензия МЧС (пожарная безопасность) — нет у поставщика"


def test_export_procurement_xlsx_subcontractors_simple_answer(
    api_client: tuple[TestClient, Path],
) -> None:
    """Отчёт: допустимость соисполнителей — простой ответ (разрешено/запрещено/
    ограничено N%), а не длинный текст условия договора."""
    import openpyxl

    client, _ = api_client

    async def _seed() -> int:
        db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
        await db.connect()
        try:
            repo = ProcurementRepository(db)
            assert await repo.upsert(
                {
                    "number": "XLSX-SUBC",
                    "platform_id": "zakupki_mos",
                    "subject": "Соисполнители ограничены",
                }
            )
            rows, _ = await repo.list_procurements(number="XLSX-SUBC")
            pid = rows[0].id
            await repo.save_requirements(
                pid,
                {
                    "subcontractors": [
                        {
                            "text": (
                                "Исполнитель вправе привлекать соисполнителей в объёме "
                                "не более 25% от цены Договора."
                            ),
                            "data": None,
                            "file_name": "dogovor.pdf",
                            "status": "limited",
                            "limit_percent": 25.0,
                        }
                    ]
                },
            )
            return pid
        finally:
            await db.dispose()

    pid = asyncio.run(_seed())
    resp = client.get(f"/api/procurements/{pid}/export.xlsx")
    assert resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert ws is not None
    rows_by_label = {row[0]: row[1] for row in ws.iter_rows(min_row=2, values_only=True)}
    assert (
        rows_by_label["Допустимость привлечения соисполнителей"]
        == "Ограничено (не более 25.0% от объёма)"
    )


def test_export_procurement_xlsx_highlights_blocking_field_mismatch(
    api_client: tuple[TestClient, Path],
) -> None:
    """Excel-экспорт (FR-13.5): отчётное LLM-поле с blocking=True и
    match=False — красным, наряду с ожидаемым значением в тексте ячейки."""
    import openpyxl

    client, _ = api_client

    async def _seed() -> int:
        db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
        await db.connect()
        try:
            repo = ProcurementRepository(db)
            assert await repo.upsert(
                {
                    "number": "XLSX-FIELD-1",
                    "platform_id": "zakupki_mos",
                    "subject": "Экспорт с блокирующим полем",
                }
            )
            rows, _ = await repo.list_procurements(number="XLSX-FIELD-1")
            pid = rows[0].id
            user = await repo.first_user()
            assert user is not None
            profile = await repo.get_active_profile(user.id)
            assert profile is not None
            await repo.update_rag_report(
                pid,
                profile.id,
                {
                    "verdict": {
                        "accepted": False,
                        "blocking_reasons": [{"source": "field:f2", "label": "объём партии"}],
                    },
                    "fields": [
                        {
                            "field_id": "f2",
                            "field_name": "объём партии",
                            "found": True,
                            "value": 300,
                            "unit": "м3",
                            "condition": {
                                "op": "llm",
                                "value_kind": "scalar",
                                "value": "не менее 500",
                            },
                            "match": False,
                            "blocking": True,
                        }
                    ],
                },
            )
            return pid
        finally:
            await db.dispose()

    pid = asyncio.run(_seed())
    resp = client.get(f"/api/procurements/{pid}/export.xlsx")
    assert resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert ws is not None
    rows_by_label = {row[0]: row[1] for row in ws.iter_rows(min_row=2, values_only=True)}
    field_value = str(rows_by_label["Поле: объём партии"])
    assert "300" in field_value
    assert "не менее 500" in field_value
    assert "НЕ выполнено" in field_value

    label_cells = {cell.value: cell for row in ws.iter_rows(min_row=2) for cell in [row[0]]}
    assert label_cells["Поле: объём партии"].font.color.rgb == "FFDC2626"


def test_procurement_requirements_post_404(api_client: tuple[TestClient, Path]) -> None:
    client, _ = api_client
    resp = client.post(
        "/api/procurements/999999/requirements",
        json={"structure": {}},
        headers=INTERNAL_HEADERS,
    )
    assert resp.status_code == 404


def test_add_procurement_by_url_reuses_existing_record(
    api_client: tuple[TestClient, Path],
) -> None:
    """US-5.5: URL уже есть в базе — данные не скачиваются повторно (живая
    подгрузка, ``fetch_procurement_by_url``, не вызывается), закупка просто
    помечается «в работе»."""
    client, _ = api_client
    url = "https://zakupki.mos.ru/need/URL-EXISTING-1"

    async def _seed() -> int:
        db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
        await db.connect()
        try:
            repo = ProcurementRepository(db)
            assert await repo.upsert(
                {
                    "number": "URL-EXISTING-1",
                    "platform_id": "zakupki_mos",
                    "subject": "Уже скачанная закупка",
                    "url": url,
                }
            )
            rows, _ = await repo.list_procurements(number="URL-EXISTING-1")
            return rows[0].id
        finally:
            await db.dispose()

    procurement_id = asyncio.run(_seed())

    resp = client.post("/api/procurements/by-url", json={"url": url})
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["id"] == procurement_id
    assert body["number"] == "URL-EXISTING-1"
    assert body["in_work"] is True


def test_add_procurement_by_url_unknown_record_live_fetch_not_a_card(
    api_client: tuple[TestClient, Path],
) -> None:
    """URL распознан по площадке (хост zakupki.mos.ru), но не похож на карточку
    закупки (нет числового needId) — 400, ничего не сохраняется."""
    client, _ = api_client
    resp = client.post(
        "/api/procurements/by-url",
        json={"url": "https://zakupki.mos.ru/purchase/list"},
    )
    assert resp.status_code == 400


def test_add_procurement_by_url_live_fetch_saves_in_work(
    api_client: tuple[TestClient, Path], monkeypatch: pytest.MonkeyPatch
) -> None:
    """Живая подгрузка по URL: запись с площадки сохраняется и сразу «в работе»."""
    client, _ = api_client
    url = "https://zakupki.gov.ru/epz/order/notice/ea20/view/common-info.html?regNumber=URL-LIVE-1"
    seen: dict[str, Any] = {}

    async def _fake_fetch(state: Any, platform_ids: list[str], fetch_url: str) -> dict[str, Any]:
        seen["platform_ids"] = platform_ids
        seen["url"] = fetch_url
        return {
            "number": "URL-LIVE-1",
            "platform_id": "zakupki_gov_44fz",
            "url": fetch_url,
            "subject": "Закупка, подгруженная по URL",
            "customer": "Заказчик по URL",
            "nmck": 1000.0,
            "is_active": True,
            "detail_json": {"number": "URL-LIVE-1"},
        }

    monkeypatch.setattr(procurements_routes, "fetch_procurement_by_url", _fake_fetch)
    resp = client.post("/api/procurements/by-url", json={"url": url})
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["number"] == "URL-LIVE-1"
    assert body["subject"] == "Закупка, подгруженная по URL"
    assert body["in_work"] is True
    # Хост ЕИС общий у 44-ФЗ и 223-ФЗ — подгрузке передаются обе площадки.
    assert set(seen["platform_ids"]) == {"zakupki_gov_44fz", "zakupki_gov_223fz"}
    assert seen["url"] == url

    # Повторное добавление того же URL — уже из базы, без живой подгрузки.
    async def _must_not_fetch(*_: Any) -> dict[str, Any]:
        raise AssertionError("живая подгрузка не должна вызываться для известного URL")

    monkeypatch.setattr(procurements_routes, "fetch_procurement_by_url", _must_not_fetch)
    again = client.post("/api/procurements/by-url", json={"url": url})
    assert again.status_code == 200, again.text
    assert again.json()["id"] == body["id"]


@pytest.mark.parametrize(
    ("error", "status"),
    [
        (ProcurementUrlError("не похоже на карточку"), 400),
        (RuntimeError("таймаут площадки"), 502),
    ],
)
def test_add_procurement_by_url_live_fetch_errors(
    api_client: tuple[TestClient, Path],
    monkeypatch: pytest.MonkeyPatch,
    error: Exception,
    status: int,
) -> None:
    """URL не похож на карточку закупки — 400 (ошибка ввода); сбой площадки —
    502. В обоих случаях ничего не сохраняется."""
    client, _ = api_client

    async def _failing_fetch(*_: Any) -> dict[str, Any]:
        raise error

    monkeypatch.setattr(procurements_routes, "fetch_procurement_by_url", _failing_fetch)
    resp = client.post(
        "/api/procurements/by-url",
        json={"url": "https://zakupki.gov.ru/epz/order/notice/ea20/view/common-info.html?x=1"},
    )
    assert resp.status_code == status, resp.text


def _seed_procurement(number: str, okpd2_codes: str | None = None) -> int:
    async def _seed() -> int:
        db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
        await db.connect()
        try:
            repo = ProcurementRepository(db)
            assert await repo.upsert(
                {
                    "number": number,
                    "platform_id": "zakupki_mos",
                    "subject": "Индексация: посев",
                    "files_json": [],
                    "okpd2_codes": okpd2_codes,
                }
            )
            rows, _ = await repo.list_procurements(number=number)
            return rows[0].id
        finally:
            await db.dispose()

    return asyncio.run(_seed())


def test_procurement_index_result_post_internal(api_client: tuple[TestClient, Path]) -> None:
    """POST /index-result (внутренний, indexing_service) — upsert в procurement_search_index."""
    client, _ = api_client
    proc_id = _seed_procurement("IDX-POST", okpd2_codes="71.20, 62.01.11")

    resp = client.post(
        f"/api/procurements/{proc_id}/index-result",
        json={
            "status": "indexed",
            "document_text": "текст технического задания",
            "content_hash": "abc123",
        },
        headers=INTERNAL_HEADERS,
    )
    assert resp.status_code == 200
    assert resp.json() == {"procurement_id": proc_id, "status": "indexed"}


def test_procurement_index_result_post_404(api_client: tuple[TestClient, Path]) -> None:
    client, _ = api_client
    resp = client.post(
        "/api/procurements/999999/index-result",
        json={"status": "indexed"},
        headers=INTERNAL_HEADERS,
    )
    assert resp.status_code == 404


def test_procurement_index_result_error_preserves_previous_document_text(
    api_client: tuple[TestClient, Path],
) -> None:
    """Повторный error-результат (транзиентный сбой) не затирает search_tsv.

    ``document_text`` больше не персистится (см. ``save_index_result``) — прежний
    успешный результат теперь проверяется через уже построенный ``search_tsv``
    (вычисляется explicit ``to_tsvector('simple', ...)`` в момент индексации).
    """
    client, _ = api_client
    proc_id = _seed_procurement("IDX-PRESERVE")

    ok = client.post(
        f"/api/procurements/{proc_id}/index-result",
        json={"status": "indexed", "document_text": "успешно извлечённый текст"},
        headers=INTERNAL_HEADERS,
    )
    assert ok.status_code == 200

    err = client.post(
        f"/api/procurements/{proc_id}/index-result",
        json={"status": "error", "error_message": "boom"},
        headers=INTERNAL_HEADERS,
    )
    assert err.status_code == 200
    assert err.json()["status"] == "error"

    async def _read_search_tsv() -> str | None:
        db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
        await db.connect()
        try:
            async with db.session() as session:
                from sqlalchemy import select

                from zakupki_parser.storage.db import ProcurementSearchIndex

                row = (
                    await session.execute(
                        select(ProcurementSearchIndex).where(
                            ProcurementSearchIndex.procurement_id == proc_id
                        )
                    )
                ).scalar_one()
                return row.search_tsv
        finally:
            await db.dispose()

    search_tsv = asyncio.run(_read_search_tsv())
    assert search_tsv is not None
    assert "извлечённый" in search_tsv


def test_index_dead_letter_requires_analyst_or_devops(
    api_client: tuple[TestClient, Path], analyst_headers: dict[str, str]
) -> None:
    """DLQ-эндпоинт (Stage A) доступен аналитику ИЛИ devops — та же политика,
    что и у остальной вкладки «Мониторинг» (см. test_monitoring_allows_analyst_
    and_devops), а не более узкая."""
    client, _ = api_client
    resp_analyst = client.get("/api/devops/index-dead-letter", headers=analyst_headers)
    assert resp_analyst.status_code == 200
    resp_devops = client.get("/api/devops/index-dead-letter")  # дефолтный клиент — admin+devops
    assert resp_devops.status_code == 200


def test_index_dead_letter_lists_entry_after_max_attempts(
    api_client: tuple[TestClient, Path],
) -> None:
    """Повторные сбои индексации (по умолчанию 5 подряд, IndexingConfig.max_attempts)
    переводят запись в dead_letter — она появляется в списке DLQ."""
    client, _ = api_client
    proc_id = _seed_procurement("IDX-DLQ-API")

    for i in range(5):
        resp = client.post(
            f"/api/procurements/{proc_id}/index-result",
            json={"status": "error", "error_message": f"boom-{i}"},
            headers=INTERNAL_HEADERS,
        )
        assert resp.status_code == 200

    dlq = client.get("/api/devops/index-dead-letter")
    assert dlq.status_code == 200
    entry = next(e for e in dlq.json()["entries"] if e["procurement_id"] == proc_id)
    assert entry["attempts"] == 5
    assert entry["error_message"] == "boom-4"


def test_index_dead_letter_retry_resets_entry(api_client: tuple[TestClient, Path]) -> None:
    """Ручной retry (аналитик/devops) сбрасывает dead-letter запись — она пропадает из DLQ."""
    client, _ = api_client
    proc_id = _seed_procurement("IDX-DLQ-RETRY")
    for i in range(5):
        client.post(
            f"/api/procurements/{proc_id}/index-result",
            json={"status": "error", "error_message": f"boom-{i}"},
            headers=INTERNAL_HEADERS,
        )
    dlq_before = client.get("/api/devops/index-dead-letter").json()["entries"]
    assert any(e["procurement_id"] == proc_id for e in dlq_before)

    resp = client.post(f"/api/devops/index-dead-letter/{proc_id}/retry")

    assert resp.status_code == 200
    assert resp.json() == {"procurement_id": proc_id, "requeued": True}
    dlq_after = client.get("/api/devops/index-dead-letter").json()["entries"]
    assert not any(e["procurement_id"] == proc_id for e in dlq_after)


def test_index_dead_letter_retry_404_for_unknown(api_client: tuple[TestClient, Path]) -> None:
    client, _ = api_client
    resp = client.post("/api/devops/index-dead-letter/999999999/retry")
    assert resp.status_code == 404


async def _upsert_platform_stats(platform_id: str, *, success: bool = True) -> None:
    """Пишет статистику одной площадки напрямую через репозиторий — у per-площадочной
    статистики нет write-эндпоинта (пишется только планировщиком), см. ``_add_region_
    procurement`` выше для того же паттерна прямой записи в тестовую БД."""
    db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
    await db.connect()
    try:
        repo = ProcurementRepository(db)
        now = datetime.now(UTC)
        await repo.upsert_platform_stats(
            platform_id=platform_id,
            iteration=1,
            started_at=now,
            finished_at=now,
            success=success,
            received=1,
            saved=1,
            error_message=None if success else "boom",
        )
    finally:
        await db.dispose()


def test_platform_stats_requires_analyst_or_devops(
    api_client: tuple[TestClient, Path], analyst_headers: dict[str, str]
) -> None:
    """Статистика по площадкам — та же политика доступа, что и у остальной вкладки
    «Мониторинг» (аналитик ИЛИ devops)."""
    client, _ = api_client
    asyncio.run(_upsert_platform_stats("api-test-platform-auth"))
    resp_analyst = client.get("/api/devops/platform-stats", headers=analyst_headers)
    assert resp_analyst.status_code == 200
    resp_devops = client.get("/api/devops/platform-stats")  # дефолтный клиент — admin+devops
    assert resp_devops.status_code == 200


def test_platform_stats_returns_items_and_shape(api_client: tuple[TestClient, Path]) -> None:
    client, _ = api_client
    asyncio.run(_upsert_platform_stats("api-test-platform-shape", success=True))

    resp = client.get("/api/devops/platform-stats", params={"search": "api-test-platform-shape"})

    assert resp.status_code == 200
    body = resp.json()
    assert body["total"] == 1
    item = body["items"][0]
    assert item["platform_id"] == "api-test-platform-shape"
    assert item["last_success"] is True
    assert item["runs_total"] == 1
    assert item["runs_failed"] == 0
    assert item["avg_received"] == pytest.approx(1.0)
    assert item["avg_saved"] == pytest.approx(1.0)


def test_platform_stats_only_failed_filter(api_client: tuple[TestClient, Path]) -> None:
    client, _ = api_client
    asyncio.run(_upsert_platform_stats("api-test-platform-ok", success=True))
    asyncio.run(_upsert_platform_stats("api-test-platform-fail", success=False))

    resp = client.get(
        "/api/devops/platform-stats",
        params={"search": "api-test-platform-", "only_failed": True},
    )

    assert resp.status_code == 200
    items = resp.json()["items"]
    assert all(i["last_success"] is False for i in items)
    assert any(i["platform_id"] == "api-test-platform-fail" for i in items)
    assert not any(i["platform_id"] == "api-test-platform-ok" for i in items)


def test_platform_stats_pagination_params(api_client: tuple[TestClient, Path]) -> None:
    client, _ = api_client
    resp = client.get(
        "/api/devops/platform-stats",
        params={"limit": 1, "offset": 0, "search": "api-test-platform"},
    )
    assert resp.status_code == 200
    assert len(resp.json()["items"]) <= 1


def test_relevance_threshold_endpoint(api_client: tuple[TestClient, Path]) -> None:
    client, _ = api_client
    body = client.get("/api/config/threshold").json()
    assert "notify_min_fit_score" in body
    assert isinstance(body["notify_min_fit_score"], (int, float))


@pytest.mark.slow
def test_list_filter_min_fit_score(api_client: tuple[TestClient, Path], inserted_id: int) -> None:
    client, _ = api_client
    # Задаём закупке фит-скор (выше порога по умолчанию 0.4).
    resp = client.post(
        f"/api/procurements/{inserted_id}/score",
        json={"profile_id": 1, "score": 123.5, "fit_score": 0.85, "score_method": "fit"},
        headers=INTERNAL_HEADERS,
    )
    assert resp.status_code == 200

    # Порог ниже/равен 0.85 — закупка попадает в выборку.
    below = client.get("/api/procurements", params={"min_fit_score": 0.5}).json()
    assert any(item["id"] == inserted_id for item in below["items"])
    assert all(
        item["fit_score"] is not None and item["fit_score"] >= 0.5 for item in below["items"]
    )

    # Порог выше 0.85 — закупка исключается.
    above = client.get("/api/procurements", params={"min_fit_score": 0.99}).json()
    assert all(item["id"] != inserted_id for item in above["items"])


def test_list_filter_min_fit_score_ignores_default_scored(
    api_client: tuple[TestClient, Path],
) -> None:
    """Дефолтный фит-скор (до обработки скорингом) не считается релевантным."""
    client, _ = api_client

    async def _insert_default() -> int:
        db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
        await db.connect()
        try:
            repo = ProcurementRepository(db)
            assert await repo.upsert(
                {
                    "number": "API-DEFAULT",
                    "platform_id": "zakupki_mos",
                    "subject": "Дефолтный скор",
                    "customer": "Заказчик ООО",
                    "fit_score": 0.9,
                    "score_method": "default",
                }
            )
            rows, _ = await repo.list_procurements(number="API-DEFAULT")
            return rows[0].id
        finally:
            await db.dispose()

    default_id = asyncio.run(_insert_default())
    asyncio.run(_match_active_profile(default_id))

    # Несмотря на высокий fit_score, дефолтный не попадает в «релевантные».
    relevant = client.get("/api/procurements", params={"min_fit_score": 0.5}).json()
    assert all(item["id"] != default_id for item in relevant["items"])
    # Но присутствует в обычном списке (без фильтра).
    all_procs = client.get("/api/procurements", params={"number": "API-DEFAULT"}).json()
    assert any(item["id"] == default_id for item in all_procs["items"])


@pytest.mark.slow
def test_sim_filtered_record_visible_with_fit_score(
    api_client: tuple[TestClient, Path],
) -> None:
    """Отсечка по векторной близости (score_method=sim) видна в API: fit_score=0.

    В «Только релевантные» (порог > 0) такая закупка не попадает, но в обычном
    списке возвращается с fit_score=0 и score_method=sim (ADR-8) — то есть
    отсечённая закупка отличима от ещё не обработанной.
    """
    client, _ = api_client

    async def _seed() -> int:
        db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
        await db.connect()
        try:
            repo = ProcurementRepository(db)
            assert await repo.upsert(
                {
                    "number": "API-VECTOR",
                    "platform_id": "zakupki_mos",
                    "subject": "Векторная отсечка",
                    "customer": "Заказчик ООО",
                }
            )
            rows, _ = await repo.list_procurements(number="API-VECTOR")
            return rows[0].id
        finally:
            await db.dispose()

    sim_id = asyncio.run(_seed())

    # Результат сервиса скоринга приходит через POST /score (ADR-7): sim —
    # предварительная фильтрация по векторной близости, LLM не выполнялся.
    resp = client.post(
        f"/api/procurements/{sim_id}/score",
        json={
            "profile_id": 1,
            "score": 0.0,
            "fit_score": 0.0,
            "score_method": "sim",
            "embedding_similarity": 0.62,
        },
        headers=INTERNAL_HEADERS,
    )
    assert resp.status_code == 200

    # В обычном списке — с fit_score=0 и score_method=sim.
    all_procs = client.get("/api/procurements", params={"number": "API-VECTOR"}).json()
    item = next(item for item in all_procs["items"] if item["id"] == sim_id)
    assert item["fit_score"] == 0.0
    assert item["score"] == 0.0
    assert item["score_method"] == "sim"
    assert item["embedding_similarity"] == 0.62

    # В «Только релевантные» (порог 0.4) не попадает: fit_score=0 < порога.
    relevant = client.get("/api/procurements", params={"min_fit_score": 0.4}).json()
    assert all(item["id"] != sim_id for item in relevant["items"])


def test_list_sort_fit_score(api_client: tuple[TestClient, Path]) -> None:
    """GET /api/procurements?sort=fit_score сортирует по релевантности (NULL в конце)."""
    client, _ = api_client

    async def _seed() -> list[int]:
        db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
        await db.connect()
        try:
            repo = ProcurementRepository(db)
            await repo.upsert(
                {
                    "number": "SORT-MID",
                    "platform_id": "zakupki_mos",
                    "subject": "Средний",
                    "fit_score": 0.5,
                    "score_method": "fit",
                }
            )
            await repo.upsert(
                {
                    "number": "SORT-NONE",
                    "platform_id": "zakupki_mos",
                    "subject": "Без скоринга",
                }
            )
            await repo.upsert(
                {
                    "number": "SORT-HIGH",
                    "platform_id": "zakupki_mos",
                    "subject": "Высокий",
                    "fit_score": 0.9,
                    "score_method": "fit",
                }
            )
            ids = []
            for num in ("SORT-MID", "SORT-NONE", "SORT-HIGH"):
                pid = await repo.find_id(num, "zakupki_mos")
                assert pid is not None
                ids.append(pid)
            return ids
        finally:
            await db.dispose()

    for pid in asyncio.run(_seed()):
        asyncio.run(_match_active_profile(pid))

    body = client.get(
        "/api/procurements",
        params={"number": "SORT", "sort": "fit_score", "limit": 100},
    ).json()
    fits = [item["fit_score"] for item in body["items"]]
    assert fits == sorted(fits, key=lambda v: v if v is not None else -1, reverse=True)
    assert body["items"][-1]["fit_score"] is None


def test_list_sort_publication_date(api_client: tuple[TestClient, Path]) -> None:
    """GET /api/procurements?sort=publication_date сортирует по убыванию даты (NULL в конце)."""
    client, _ = api_client

    async def _seed() -> list[int]:
        db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
        await db.connect()
        try:
            repo = ProcurementRepository(db)
            await repo.upsert(
                {
                    "number": "SORTDATE-OLD",
                    "platform_id": "zakupki_mos",
                    "subject": "Старая",
                    "publication_date": datetime(2026, 1, 1, tzinfo=UTC),
                }
            )
            await repo.upsert(
                {
                    "number": "SORTDATE-NEW",
                    "platform_id": "zakupki_mos",
                    "subject": "Новая",
                    "publication_date": datetime(2026, 6, 1, tzinfo=UTC),
                }
            )
            await repo.upsert(
                {
                    "number": "SORTDATE-NONE",
                    "platform_id": "zakupki_mos",
                    "subject": "Без даты",
                }
            )
            ids = []
            for num in ("SORTDATE-OLD", "SORTDATE-NEW", "SORTDATE-NONE"):
                pid = await repo.find_id(num, "zakupki_mos")
                assert pid is not None
                ids.append(pid)
            return ids
        finally:
            await db.dispose()

    for pid in asyncio.run(_seed()):
        asyncio.run(_match_active_profile(pid))

    body = client.get(
        "/api/procurements",
        params={"number": "SORTDATE", "sort": "publication_date", "limit": 100},
    ).json()
    dates = [item["publication_date"] for item in body["items"]]
    assert dates[0] == "2026-06-01T00:00:00Z"
    assert dates[1] == "2026-01-01T00:00:00Z"
    assert dates[-1] is None


@pytest.mark.slow
def test_set_score_by_external_service(
    api_client: tuple[TestClient, Path], inserted_id: int
) -> None:
    client, _ = api_client
    resp = client.post(
        f"/api/procurements/{inserted_id}/score",
        json={"profile_id": 1, "score": 123.5, "fit_score": 0.85, "score_method": "fit"},
        headers=INTERNAL_HEADERS,
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["score"] == 123.5
    assert body["fit_score"] == 0.85
    assert body["score_method"] == "fit"

    detail = client.get(f"/api/procurements/{inserted_id}").json()
    assert detail["score"] == 123.5
    assert detail["fit_score"] == 0.85


@pytest.mark.slow
def test_set_score_notifies_above_threshold(
    api_client: tuple[TestClient, Path], inserted_id: int
) -> None:
    client, _ = api_client
    calls: list[dict[str, object]] = []

    class _FakeNotifier:
        async def notify(self, record: dict[str, object]) -> None:
            calls.append(record)

    state = cast(Any, client.app).state.parser
    state.notifier = _FakeNotifier()
    state.notify_min_fit_score = 0.5

    # Ниже порога — score обновляется, уведомления нет (sim — терминальная
    # отсечка по векторной близости, ADR-8).
    resp = client.post(
        f"/api/procurements/{inserted_id}/score",
        json={"profile_id": 1, "score": 50.0, "fit_score": 0.3, "score_method": "sim"},
        headers=INTERNAL_HEADERS,
    )
    assert resp.status_code == 200
    assert calls == []

    # Выше порога (по fit_score) — уведомление с обновлённой карточкой.
    resp = client.post(
        f"/api/procurements/{inserted_id}/score",
        json={"profile_id": 1, "score": 150.0, "fit_score": 0.9, "score_method": "fit"},
        headers=INTERNAL_HEADERS,
    )
    assert resp.status_code == 200
    assert len(calls) == 1
    assert calls[0]["score"] == 150.0
    assert calls[0]["fit_score"] == 0.9


def test_set_score_404(api_client: tuple[TestClient, Path]) -> None:
    client, _ = api_client
    assert (
        client.post(
            "/api/procurements/999999/score",
            json={"profile_id": 1, "score": 1.0, "score_method": "fit"},
            headers=INTERNAL_HEADERS,
        ).status_code
        == 404
    )


def test_set_score_rejects_unknown_method(
    api_client: tuple[TestClient, Path], inserted_id: int
) -> None:
    """POST /score с неизвестным score_method отклоняется (422), а не пишется в БД.

    Приёмный эндпоинт принимает только известные результаты внешнего скоринга
    (fit/pwin/margin/sim, ADR-7/ADR-8): неизвестный метод — признак рассинхрона
    конвейера, его не нужно молча сохранять.
    """
    client, _ = api_client
    resp = client.post(
        f"/api/procurements/{inserted_id}/score",
        json={"profile_id": 1, "score": 50.0, "fit_score": 0.3, "score_method": "unknown-stage"},
        headers=INTERNAL_HEADERS,
    )
    assert resp.status_code == 422

    detail = client.get(f"/api/procurements/{inserted_id}").json()
    assert detail["score_method"] != "unknown-stage"


def test_procurement_has_customer_id_and_name(
    api_client: tuple[TestClient, Path], inserted_id: int
) -> None:
    client, _ = api_client
    body = client.get(f"/api/procurements/{inserted_id}").json()
    assert body["customer_id"] is not None
    assert body["customer"] == "Заказчик ООО"


def test_procurements_filter_by_customer(
    api_client: tuple[TestClient, Path], inserted_id: int
) -> None:
    client, _ = api_client
    resp = client.get("/api/procurements", params={"customer": "заказчик"})
    assert resp.status_code == 200
    body = resp.json()
    assert body["total"] >= 1
    assert any(item["id"] == inserted_id for item in body["items"])


def test_procurement_region_out_and_filter(
    api_client: tuple[TestClient, Path], inserted_id: int
) -> None:
    """Регион в карточке закупки и серверный фильтр списка по региону."""
    client, _ = api_client
    # inserted_id создан без региона (region=None в ответе).
    detail = client.get(f"/api/procurements/{inserted_id}").json()
    assert "region" in detail
    assert detail["region"] is None

    asyncio.run(_add_region_procurement())
    resp = client.get("/api/procurements", params={"region": "Москва"})
    assert resp.status_code == 200
    body = resp.json()
    assert body["total"] >= 1
    hit = next(item for item in body["items"] if item["region"] == "г. Москва")
    assert hit["number"] == "RG-API-1"
    assert hit["region"] == "г. Москва"


def test_add_exclusion_word_endpoint(api_client: tuple[TestClient, Path]) -> None:
    """POST /api/procurements/{id}/exclusion-word — кнопка «В исключения» карточки:
    добавляет фразу в исключения активного профиля без отбраковки закупки."""
    client, _ = api_client
    pid = _seed_procurement("EXCL-API-1")

    resp = client.post(f"/api/procurements/{pid}/exclusion-word", json={"word": "не наш профиль*"})
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["added"] is True
    assert "не наш профиль*" in body["exclusion_words"]

    # Повторное добавление той же фразы — идемпотентно (added=False).
    resp2 = client.post(f"/api/procurements/{pid}/exclusion-word", json={"word": "не наш профиль*"})
    assert resp2.status_code == 200
    assert resp2.json()["added"] is False

    # Закупка не отклонена — только по прямой отбраковке (reject).
    detail = client.get(f"/api/procurements/{pid}").json()
    assert detail["score_method"] != "reject"


def test_add_exclusion_word_rejects_blank(api_client: tuple[TestClient, Path]) -> None:
    client, _ = api_client
    pid = _seed_procurement("EXCL-API-2")
    resp = client.post(f"/api/procurements/{pid}/exclusion-word", json={"word": "   "})
    assert resp.status_code == 422


def test_add_exclusion_word_404_for_unknown_procurement(
    api_client: tuple[TestClient, Path],
) -> None:
    client, _ = api_client
    resp = client.post("/api/procurements/999999/exclusion-word", json={"word": "x"})
    assert resp.status_code == 404


def test_set_score_region_mismatch_not_written(api_client: tuple[TestClient, Path]) -> None:
    """Регион вне целевых профиля (стал известен после досборки) — скор не пишется.

    Повторная клиентская фильтрация в POST /score (BR-08): результат отклоняется,
    оценка профиля удаляется — профиль «не отобрал» закупку.
    """
    client, _ = api_client
    pid = asyncio.run(_add_region_score_procurement("RG-SCORE-1", "Санкт-Петербург"))
    created = client.post(
        "/api/clients",
        json={
            "name": "region-check",
            "competencies": COMP_JSON,
            "target_regions": ["Москва"],
        },
    )
    assert created.status_code == 200
    profile_id = created.json()["id"]

    resp = client.post(
        f"/api/procurements/{pid}/score",
        json={"profile_id": profile_id, "score": 42.0, "fit_score": 0.9, "score_method": "fit"},
        headers=INTERNAL_HEADERS,
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["score"] is None
    assert body["fit_score"] is None
    assert asyncio.run(_has_evaluation(pid, profile_id)) is False


def test_set_score_region_match_written(api_client: tuple[TestClient, Path]) -> None:
    """Регион совпал с целевым — результат скоринга записывается как обычно."""
    client, _ = api_client
    pid = asyncio.run(_add_region_score_procurement("RG-SCORE-2", "г. Москва"))
    created = client.post(
        "/api/clients",
        json={
            "name": "region-match",
            "competencies": COMP_JSON,
            "target_regions": ["Москва"],
        },
    )
    assert created.status_code == 200
    profile_id = created.json()["id"]

    resp = client.post(
        f"/api/procurements/{pid}/score",
        json={"profile_id": profile_id, "score": 77.0, "fit_score": 0.88, "score_method": "fit"},
        headers=INTERNAL_HEADERS,
    )
    assert resp.status_code == 200
    assert resp.json()["fit_score"] == 0.88
    assert asyncio.run(_has_evaluation(pid, profile_id)) is True


def test_set_score_too_far_not_written(api_client: tuple[TestClient, Path]) -> None:
    """Вердикт analysis_service «дальше max_region_distance_km» — скор не пишется.

    Вердикт (rag_report['geo']['too_far']) считает analysis_service (координат
    у парсера нет); здесь он применяется как ещё одна клиентская пост-фильтрация,
    симметрично проверке региона — оценка профиля удаляется.
    """
    client, _ = api_client
    pid = asyncio.run(_add_region_score_procurement("RG-SCORE-3", "Новосибирская область"))
    created = client.post(
        "/api/clients",
        json={
            "name": "region-distance",
            "competencies": COMP_JSON,
            "target_regions": ["Москва"],
            "max_region_distance_km": 50,
        },
    )
    assert created.status_code == 200
    profile_id = created.json()["id"]

    resp = client.post(
        f"/api/procurements/{pid}/score",
        json={
            "profile_id": profile_id,
            "score": 42.0,
            "fit_score": 0.9,
            "score_method": "fit",
            "rag_report": {
                "geo": {
                    "too_far": True,
                    "distance_km": 3200.0,
                    "max_distance_km": 50.0,
                }
            },
        },
        headers=INTERNAL_HEADERS,
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["score"] is None
    assert body["fit_score"] is None
    assert asyncio.run(_has_evaluation(pid, profile_id)) is False


def test_set_score_within_distance_written(api_client: tuple[TestClient, Path]) -> None:
    """Вердикт analysis_service «в пределах max_region_distance_km» — скор пишется как обычно."""
    client, _ = api_client
    pid = asyncio.run(_add_region_score_procurement("RG-SCORE-4", "Москва"))
    created = client.post(
        "/api/clients",
        json={
            "name": "region-distance-ok",
            "competencies": COMP_JSON,
            "target_regions": ["Москва"],
            "max_region_distance_km": 50,
        },
    )
    assert created.status_code == 200
    profile_id = created.json()["id"]

    resp = client.post(
        f"/api/procurements/{pid}/score",
        json={
            "profile_id": profile_id,
            "score": 55.0,
            "fit_score": 0.7,
            "score_method": "fit",
            "rag_report": {
                "geo": {
                    "too_far": False,
                    "distance_km": 12.0,
                    "max_distance_km": 50.0,
                }
            },
        },
        headers=INTERNAL_HEADERS,
    )
    assert resp.status_code == 200
    assert resp.json()["fit_score"] == 0.7
    assert asyncio.run(_has_evaluation(pid, profile_id)) is True


def test_profile_from_url_rejects_non_http_scheme(
    api_client: tuple[TestClient, Path],
) -> None:
    """Маршрут смонтирован и требует авторизации (клиент фикстуры аутентифицирован);
    некорректная схема URL отклоняется до сети — понятная ошибка 400, а не 500."""
    client, _ = api_client
    resp = client.post(
        "/api/clients/profile/from-url",
        json={"url": "file:///etc/passwd"},
    )
    assert resp.status_code == 400
    assert "http/https" in resp.json()["detail"]


def test_profile_website_url_persists_across_save(api_client: tuple[TestClient, Path]) -> None:
    """Сайт поставщика (кнопка «Заполнить профиль по URL») сохраняется вместе
    с профилем — не нужно вводить его заново при повторном заполнении."""
    client, _ = api_client
    created = client.post(
        "/api/clients",
        json={
            "name": "website-url-profile",
            "competencies": COMP_JSON,
            "website_url": "https://example.com",
        },
    )
    assert created.status_code == 200, created.text
    profile_id = created.json()["id"]
    assert created.json()["website_url"] == "https://example.com"

    fetched = client.get(f"/api/clients/{profile_id}")
    assert fetched.json()["website_url"] == "https://example.com"

    cleared = client.put(
        f"/api/clients/{profile_id}",
        json={"name": "website-url-profile", "competencies": COMP_JSON, "website_url": None},
    )
    assert cleared.status_code == 200, cleared.text
    assert cleared.json()["website_url"] is None


def test_profile_report_fields_roundtrip(api_client: tuple[TestClient, Path]) -> None:
    """Конструктор отчётных полей (FR-12.1): поля профиля сохраняются полной
    заменой вместе с профилем и возвращаются в каноническом виде
    (``normalize_report_fields``); условие и блокировка — часть того же поля
    профиля (FR-13.5, FR-13.7)."""
    from scoring_common.conditions import normalize_report_fields

    client, _ = api_client
    fields = [
        {
            "id": "f1",
            "name": "код ФККО",
            "hint": "код отхода по ФККО",
            "type": "string",
            "unit": None,
        },
        {
            "id": "f2",
            "name": "объём партии",
            "hint": "объём вывоза",
            "type": "number",
            "unit": "м3",
            "condition": {"op": "gte", "value": "500"},
            "blocking": True,
        },
    ]
    created = client.post(
        "/api/clients",
        json={"name": "report-fields-profile", "competencies": COMP_JSON, "report_fields": fields},
    )
    assert created.status_code == 200, created.text
    profile_id = created.json()["id"]
    expected = normalize_report_fields(fields)
    assert expected[1]["condition"] == {"op": "gte", "value_kind": "scalar", "value": "500"}
    assert created.json()["report_fields"] == expected

    fetched = client.get(f"/api/clients/{profile_id}")
    assert fetched.json()["report_fields"] == expected

    replaced = client.put(
        f"/api/clients/{profile_id}",
        json={
            "name": "report-fields-profile",
            "competencies": COMP_JSON,
            "report_fields": [fields[0]],
        },
    )
    assert replaced.status_code == 200, replaced.text
    assert replaced.json()["report_fields"] == [expected[0]]


def test_customers_list_and_rating(api_client: tuple[TestClient, Path], inserted_id: int) -> None:
    client, _ = api_client
    customer_id = client.get(f"/api/procurements/{inserted_id}").json()["customer_id"]

    # /api/customers сужен до заказчиков, связанных с закупками активного
    # профиля вызывающего (BR-07) — inserted_id уже отобрана дефолтным
    # профилем фикстурой ``inserted_id`` (иначе не попала бы и в список
    # закупок, см. list_procurements).
    listed = client.get("/api/customers").json()
    assert listed["total"] >= 1
    assert any(item["id"] == customer_id for item in listed["items"])

    got = client.get(f"/api/customers/{customer_id}")
    assert got.status_code == 200
    assert got.json()["name"] == "Заказчик ООО"

    rated = client.post(
        f"/api/customers/{customer_id}/rating", json={"rating": 0.9}, headers=INTERNAL_HEADERS
    )
    assert rated.status_code == 200
    assert rated.json()["rating"] == 0.9


def test_customers_list_excludes_customer_not_matched_by_active_profile(
    api_client: tuple[TestClient, Path],
) -> None:
    """Заказчик закупки, не отобранной активным профилем вызывающего (нет
    записи в procurement_evaluations для него), не попадает в /api/customers —
    даже если закупка сохранена в общей БД (закупки общие для всех профилей)."""
    client, _ = api_client

    async def _seed_unmatched() -> int:
        db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
        await db.connect()
        try:
            repo = ProcurementRepository(db)
            await repo.upsert(
                {
                    "number": "CUST-UNMATCHED-1",
                    "platform_id": "zakupki_mos",
                    "subject": "Не отобрано профилем",
                    "customer": "ООО Не в этом профиле",
                }
            )
            pid = await repo.find_id("CUST-UNMATCHED-1", "zakupki_mos")
            assert pid is not None
            return pid
        finally:
            await db.dispose()

    asyncio.run(_seed_unmatched())

    listed = client.get("/api/customers", params={"limit": 100}).json()
    assert not any(item["name"] == "ООО Не в этом профиле" for item in listed["items"])


def test_customer_rating_404(api_client: tuple[TestClient, Path]) -> None:
    client, _ = api_client
    assert (
        client.post(
            "/api/customers/999999/rating", json={"rating": 1.0}, headers=INTERNAL_HEADERS
        ).status_code
        == 404
    )
    assert client.get("/api/customers/999999").status_code == 404


def test_config_get_redacts_and_put_saves(tmp_path: Path, analyst_headers: dict[str, str]) -> None:
    """Конфиг-сервис: GET отдаёт без секретов, PUT валидирует и пишет в YAML.

    Используем копию configs в tmp_path, чтобы не трогать реальный конфиг.
    """
    from zakupki_parser.api.app import create_app

    cfgdir = tmp_path / "configs"
    # Копируем ТЕСТОВЫЙ набор конфигов (tests/configs), а не рабочие configs/*.
    shutil.copytree(Path(__file__).resolve().parents[2] / "tests" / "configs", cfgdir)
    os.environ["ZAKUPKI_DB_DSN"] = TEST_DSN
    app = create_app(str(cfgdir))
    with TestClient(app) as client:
        client.headers.update(analyst_headers)
        cfg = client.get("/api/config").json()
        assert "sites" in cfg
        # Эксплуатационные параметры (таймер, БД, уведомления) не отдаются через API —
        # они живут в config_ops.yaml.
        assert "timeout_seconds" not in cfg
        assert "notifications" not in cfg

        old = cfg["default_cutoff_days"]
        cfg["default_cutoff_days"] = old + 1
        r = client.put("/api/config", json=cfg)
        assert r.status_code == 200
        assert r.json()["default_cutoff_days"] == old + 1

        saved = (cfgdir / "config_service.yaml").read_text(encoding="utf-8")
        assert f"default_cutoff_days: {old + 1}" in saved

        # Некорректные данные — 422, файл не меняется.
        bad = client.put("/api/config", json={"default_cutoff_days": "not-a-number"})
        assert bad.status_code == 422
    os.environ.pop("ZAKUPKI_DB_DSN", None)


def test_devops_indexing_config_get_put_and_role_gate(tmp_path: Path) -> None:
    """Devops-редактор фоновой индексации: GET/PUT только indexing:, без секретов
    и без затирания остальных полей config_service.yaml (sites/scoring)."""
    from zakupki_parser.api.app import create_app

    cfgdir = tmp_path / "configs"
    shutil.copytree(Path(__file__).resolve().parents[2] / "tests" / "configs", cfgdir)
    os.environ["ZAKUPKI_DB_DSN"] = TEST_DSN
    app = create_app(str(cfgdir))
    with TestClient(app) as client:

        async def _seed_users() -> tuple[int, int]:
            db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
            await db.connect()
            try:
                repo = ProcurementRepository(db)
                devops_user = await repo.create_user("devops-idx", "h", [ROLE_DEVOPS])
                analyst_user = await repo.create_user("analyst-idx", "h", [ROLE_ANALYST, ROLE_USER])
                return devops_user.id, analyst_user.id
            finally:
                await db.dispose()

        devops_id, analyst_id = asyncio.run(_seed_users())
        devops_headers = {
            "Authorization": f"Bearer {create_token(devops_id, [ROLE_DEVOPS], AUTH_SECRET, 3600)}"
        }
        analyst_headers_local = {
            "Authorization": (
                f"Bearer {create_token(analyst_id, [ROLE_ANALYST, ROLE_USER], AUTH_SECRET, 3600)}"
            )
        }

        # Аналитик не может дёргать devops-эндпоинт (403), даже зная его.
        r = client.get("/api/devops/indexing-config", headers=analyst_headers_local)
        assert r.status_code == 403

        cfg = client.get("/api/devops/indexing-config", headers=devops_headers).json()
        assert "enabled" in cfg and "okpd2_prefixes" in cfg and "excluded_platforms" in cfg

        r = client.put(
            "/api/devops/indexing-config",
            json={"enabled": True, "okpd2_prefixes": ["62.2"], "excluded_platforms": []},
            headers=devops_headers,
        )
        assert r.status_code == 200
        assert r.json()["okpd2_prefixes"] == ["62.2"]

        saved = (cfgdir / "config_service.yaml").read_text(encoding="utf-8")
        assert "62.2" in saved
        assert "sites:" in saved  # остальные поля не затёрты

        bad = client.put(
            "/api/devops/indexing-config",
            json={"okpd2_prefixes": "not-a-list"},
            headers=devops_headers,
        )
        assert bad.status_code == 422
    os.environ.pop("ZAKUPKI_DB_DSN", None)


def test_export_csv_download(api_client: tuple[TestClient, Path], inserted_id: int) -> None:
    """CSV отдаётся файлом: только активные релевантные закупки (fit_score >= порога)."""
    client, _ = api_client

    async def _seed_relevant() -> None:
        db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
        await db.connect()
        try:
            repo = ProcurementRepository(db)
            await repo.upsert(
                {
                    "number": "EXPORT-REL",
                    "platform_id": "zakupki_mos",
                    "subject": "Релевантная активная",
                    "customer": "Заказчик ООО",
                }
            )
            await repo.upsert(
                {
                    "number": "EXPORT-IRR",
                    "platform_id": "zakupki_mos",
                    "subject": "Нерелевантная",
                }
            )
            await repo.upsert(
                {
                    "number": "EXPORT-INACTIVE",
                    "platform_id": "zakupki_mos",
                    "subject": "Неактивная",
                    "is_active": False,
                }
            )
            rows, _ = await repo.list_procurements(number="EXPORT-")
            ids = {p.number: p.id for p in rows}
            # Оценки per-profile (BR-07): релевантность фильтра/выгрузки считается
            # по procurement_evaluations активного профиля.
            user = await repo.first_user()
            assert user is not None
            profile = await repo.get_active_profile(user.id)
            assert profile is not None
            await repo.upsert_score(
                ids["EXPORT-REL"], profile.id, fit_score=0.8, score_method="fit"
            )
            await repo.upsert_score(
                ids["EXPORT-IRR"], profile.id, fit_score=0.2, score_method="fit"
            )
        finally:
            await db.dispose()

    asyncio.run(_seed_relevant())

    resp = client.post("/api/procurements/export", json={"min_fit_score": 0.4})
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/csv")
    assert "attachment" in resp.headers.get("content-disposition", "")
    content = resp.content.decode("utf-8-sig")
    # Заголовок + активная релевантная запись.
    assert "number,platform_id" in content
    assert "EXPORT-REL" in content
    assert "Заказчик ООО" in content
    # Нерелевантная (fit_score < 0.4) и неактивная закупки в выгрузку не попадают.
    assert "EXPORT-IRR" not in content
    assert "EXPORT-INACTIVE" not in content


def test_prompts_list_get_put_validate(tmp_path: Path, analyst_headers: dict[str, str]) -> None:
    """Промпты: список, чтение, сохранение; JSON валидируется, traversal запрещён.

    Используем копию tests/configs и отдельный каталог промптов в tmp_path,
    чтобы не трогать реальные конфиги и файлы промптов.
    """
    from zakupki_parser.api.app import create_app

    cfgdir = tmp_path / "configs"
    shutil.copytree(Path(__file__).resolve().parents[2] / "tests" / "configs", cfgdir)
    prompts_dir = tmp_path / "prompts"
    prompts_dir.mkdir()
    (prompts_dir / "fit_system.md").write_text("СТАРЫЙ ПРОМПТ", encoding="utf-8")
    (prompts_dir / "few_shot.json").write_text('[{"a": 1}]', encoding="utf-8")

    os.environ["ZAKUPKI_DB_DSN"] = TEST_DSN
    os.environ["ZAKUPKI_PROMPTS_DIR"] = str(prompts_dir)
    app = create_app(str(cfgdir))
    with TestClient(app) as client:
        client.headers.update(analyst_headers)
        # Список: только md/json внутри prompts_dir.
        files = client.get("/api/prompts").json()["files"]
        names = [f["name"] for f in files]
        assert "fit_system.md" in names
        assert "few_shot.json" in names

        # Чтение содержимого.
        got = client.get("/api/prompts/fit_system.md")
        assert got.status_code == 200
        assert got.json()["content"] == "СТАРЫЙ ПРОМПТ"
        assert got.json()["kind"] == "markdown"

        # Сохранение markdown.
        r = client.put("/api/prompts/fit_system.md", json={"content": "НОВЫЙ ПРОМПТ"})
        assert r.status_code == 200
        assert r.json()["content"] == "НОВЫЙ ПРОМПТ"
        assert (prompts_dir / "fit_system.md").read_text(encoding="utf-8") == "НОВЫЙ ПРОМПТ"

        # Некорректный JSON — 422, файл не меняется.
        bad = client.put("/api/prompts/few_shot.json", json={"content": "{broken"})
        assert bad.status_code == 422
        assert (prompts_dir / "few_shot.json").read_text(encoding="utf-8") == '[{"a": 1}]'

        # Корректный JSON сохраняется.
        ok = client.put("/api/prompts/few_shot.json", json={"content": '[{"b": 2}]'})
        assert ok.status_code == 200
        assert ok.json()["kind"] == "json"

        # Path traversal и несуществующие файлы отклоняются.
        assert client.get("/api/prompts/..%2Fconfig_service.yaml").status_code in (400, 404)
        assert client.put("/api/prompts/nope.md", json={"content": "x"}).status_code == 404
        assert client.get("/api/prompts/secret.txt").status_code in (400, 404)
    os.environ.pop("ZAKUPKI_DB_DSN", None)
    os.environ.pop("ZAKUPKI_PROMPTS_DIR", None)


def test_analysis_prompts_list_get_put(tmp_path: Path, analyst_headers: dict[str, str]) -> None:
    """Промпты analysis_service: список, чтение, сохранение.

    Используем отдельный каталог промптов анализатора в tmp_path (env
    ZAKUPKI_ANALYSIS_PROMPTS_DIR), чтобы не трогать реальные файлы.
    """
    from zakupki_parser.api.app import create_app

    cfgdir = tmp_path / "configs"
    shutil.copytree(Path(__file__).resolve().parents[2] / "tests" / "configs", cfgdir)
    analysis_prompts_dir = tmp_path / "analysis_prompts"
    analysis_prompts_dir.mkdir()
    (analysis_prompts_dir / "verdict_system.md").write_text("СТАРЫЙ ПРОМПТ", encoding="utf-8")

    os.environ["ZAKUPKI_DB_DSN"] = TEST_DSN
    os.environ["ZAKUPKI_ANALYSIS_PROMPTS_DIR"] = str(analysis_prompts_dir)
    app = create_app(str(cfgdir))
    with TestClient(app) as client:
        client.headers.update(analyst_headers)
        # Список: только md/json внутри каталога промптов анализатора.
        files = client.get("/api/analysis-prompts").json()["files"]
        names = [f["name"] for f in files]
        assert "verdict_system.md" in names

        # Чтение содержимого.
        got = client.get("/api/analysis-prompts/verdict_system.md")
        assert got.status_code == 200
        assert got.json()["content"] == "СТАРЫЙ ПРОМПТ"
        assert got.json()["kind"] == "markdown"

        # Сохранение.
        r = client.put("/api/analysis-prompts/verdict_system.md", json={"content": "НОВЫЙ ПРОМПТ"})
        assert r.status_code == 200
        assert r.json()["content"] == "НОВЫЙ ПРОМПТ"
        assert (analysis_prompts_dir / "verdict_system.md").read_text(
            encoding="utf-8"
        ) == "НОВЫЙ ПРОМПТ"

        # Каталоги скоринга и анализатора независимы: анализ-промпт не виден
        # в /api/prompts и наоборот.
        assert client.get("/api/analysis-prompts/fit_system.md").status_code == 404

        # Path traversal отклоняется.
        assert client.get("/api/analysis-prompts/..%2Fconfig_service.yaml").status_code in (
            400,
            404,
        )
    os.environ.pop("ZAKUPKI_DB_DSN", None)
    os.environ.pop("ZAKUPKI_ANALYSIS_PROMPTS_DIR", None)


def test_active_context_without_profile_is_conflict(
    api_client: tuple[TestClient, Path],
) -> None:
    """Пользователь без профилей получает 409 — профиль не досоздаётся на лету
    (профиль создаётся вместе с пользователем и при выдаче роли user/analyst)."""
    client, _ = api_client

    async def _mk_user() -> int:
        db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
        await db.connect()
        try:
            repo = ProcurementRepository(db)
            user = await repo.create_user("no-profile-user", "h", [ROLE_USER])
            return user.id
        finally:
            await db.dispose()

    user_id = asyncio.run(_mk_user())
    token = create_token(user_id, [ROLE_USER], AUTH_SECRET, 3600)
    resp = client.get("/api/procurements", headers={"Authorization": f"Bearer {token}"})
    assert resp.status_code == 409


def test_active_context_uses_disabled_profile(
    api_client: tuple[TestClient, Path],
) -> None:
    """Выключенный профиль выбирается активным контекстом (FR-1.3).

    Активность профиля не зависит от ``enabled``: отключённый от постоянного
    мониторинга профиль можно использовать для ручной обработки «в работе».
    Если у пользователя есть хотя бы один профиль, активный контекст разрешается
    всегда (первый по id), 409 не возвращается.
    """
    client, _ = api_client

    async def _mk_user() -> int:
        db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
        await db.connect()
        try:
            repo = ProcurementRepository(db)
            user = await repo.create_user("disabled-profile-user", "h", [ROLE_USER])
            await repo.upsert_profile(
                {
                    "name": "disabled-only",
                    "enabled": False,
                    "is_active": False,
                    "competencies": COMP_JSON,
                    "keywords": [],
                    "exclusion_words": [],
                },
                user.id,
            )
            return user.id
        finally:
            await db.dispose()

    user_id = asyncio.run(_mk_user())
    token = create_token(user_id, [ROLE_USER], AUTH_SECRET, 3600)
    headers = {"Authorization": f"Bearer {token}"}

    resp = client.get("/api/procurements", headers=headers)
    assert resp.status_code == 200

    async def _active_is_disabled_profile() -> None:
        db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
        await db.connect()
        try:
            repo = ProcurementRepository(db)
            rows, total = await repo.list_profiles(user_id)
            assert total == 1
            assert rows[0].name == "disabled-only"
            active = await repo.get_active_profile(user_id)
            assert active is not None
            assert active.name == "disabled-only"
        finally:
            await db.dispose()

    asyncio.run(_active_is_disabled_profile())


def test_monitoring_allows_analyst_and_devops(
    api_client: tuple[TestClient, Path], analyst_headers: dict[str, str]
) -> None:
    """Вкладка «Мониторинг» — аналитик ИЛИ devops (изначально была devops-only,
    см. plan «radiant-crunching-lemon»; расширена под Dead Letter Queue фоновой
    индексации — Stage A плана «индекс как основной механизм discovery», ничего
    из состава ответа не чувствительно для аналитика).

    ``analyst_headers`` — токен реального пользователя с ролями analyst/user (без
    devops) в БД: ``require_user`` перечитывает роли из БД по ``sub`` токена, так
    что важен фактический набор ролей пользователя, а не то, что записано в payload.
    """
    client, _ = api_client
    resp = client.get("/api/devops/monitoring", headers=analyst_headers)
    assert resp.status_code == 200


def test_monitoring_returns_queues_index_and_resources(
    api_client: tuple[TestClient, Path],
) -> None:
    """Без настроенного scoring_transport — best-effort «недоступен», а не 500."""
    client, _ = api_client
    resp = client.get("/api/devops/monitoring")
    assert resp.status_code == 200
    body = resp.json()
    assert body["queues"] == {"available": False}
    # ``enabled``/``okpd2_prefixes`` зависят от текущего configs/config_service.yaml
    # (не фиксируем конкретное значение — только форму ответа).
    assert isinstance(body["index"]["enabled"], bool)
    assert isinstance(body["index"]["counts"], dict)
    # Не фиксируем пустой список: api_client — module-scoped, другие тесты в этом
    # файле (например, test_procurement_index_result_error_preserves_previous_
    # document_text) могли оставить свои status='error' строки в общей БД.
    assert isinstance(body["index"]["recent_errors"], list)
    assert 0.0 <= body["resources"]["memory"]["percent"] <= 100.0
    assert body["resources"]["disk"]["total"] > 0
    # processes: как минимум сам процесс API-теста — форма ответа (доля с момента
    # предыдущего опроса, не фиксированное окно, поэтому конкретные cpu_percent не
    # проверяем).
    processes = body["resources"]["processes"]
    assert isinstance(processes, list)
    assert len(processes) >= 1
    expected_keys = {"pid", "label", "cpu_percent", "rss_bytes", "rss_percent"}
    assert all(expected_keys <= p.keys() for p in processes)
    # cycles/storage — форма ответа (значения зависят от истории проходов/локальной
    # ФС окружения, где запущены тесты — не фиксируем конкретные числа).
    assert "last" in body["cycles"]
    assert "average" in body["cycles"]
    assert isinstance(body["storage"]["file_storage_bytes"], int)
    assert isinstance(body["storage"]["db_bytes"], int)
    assert body["storage"]["db_bytes"] > 0


def test_profile_report_field_invalid_condition_rejected(
    api_client: tuple[TestClient, Path],
) -> None:
    """Условие проверяется при записи профиля: оператор не подходит к типу поля."""
    client, _ = api_client
    resp = client.post(
        "/api/clients",
        json={
            "name": "bad-condition-profile",
            "competencies": COMP_JSON,
            "report_fields": [
                {
                    "id": "f1",
                    "name": "объём",
                    "type": "number",
                    "condition": {"op": "all_in", "value": ["1"]},
                }
            ],
        },
    )
    assert resp.status_code == 400
    assert "объём" in resp.json()["detail"]


def test_condition_change_rechecks_reports_without_llm(
    api_client: tuple[TestClient, Path],
) -> None:
    """Правка условия поля пересчитывает сохранённые отчёты профиля без LLM:
    match и авто-отклонение меняются, отчёт остаётся актуальным (кнопка
    «Анализ» не загорается), ход пересчёта виден в /recheck."""
    from scoring_common.conditions import apply_condition, extraction_key, normalize_report_fields

    client, _ = api_client
    active = client.get("/api/clients/active").json()
    profile_id = active["id"]

    def _field(threshold: str) -> dict[str, Any]:
        return {
            "id": "vol",
            "name": "объём партии",
            "type": "number",
            "unit": "м3",
            "condition": {"op": "gte", "value": threshold},
            "blocking": True,
        }

    def _put(threshold: str) -> dict[str, Any]:
        resp = client.put(
            f"/api/clients/{profile_id}",
            json={
                "name": active["name"],
                "competencies": COMP_JSON,
                "report_fields": [_field(threshold)],
            },
        )
        assert resp.status_code == 200, resp.text
        return cast(dict[str, Any], resp.json())

    def _wait_recheck() -> dict[str, Any]:
        for _ in range(100):
            status = cast(dict[str, Any], client.get(f"/api/clients/{profile_id}/recheck").json())
            if not status["running"]:
                return status
            time.sleep(0.05)
        raise AssertionError("пересчёт условий не завершился")

    saved = _put("500")
    _wait_recheck()
    field_def = normalize_report_fields([_field("500")])[0]
    stored = apply_condition(
        {
            "field_id": "vol",
            "field_name": "объём партии",
            "field_type": "number",
            "unit": "м3",
            "found": True,
            "value": 300.0,
            "extraction_key": extraction_key(field_def),
        },
        field_def,
    )
    assert stored["match"] is False

    async def _seed() -> int:
        db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
        await db.connect()
        try:
            repo = ProcurementRepository(db)
            assert await repo.upsert(
                {"number": "RECHECK-1", "platform_id": "zakupki_mos", "subject": "Пересчёт"}
            )
            rows, _ = await repo.list_procurements(number="RECHECK-1")
            pid = rows[0].id
            await repo.upsert_score(
                pid,
                profile_id,
                score_method="fit",
                rag_report={
                    "status": "ok",
                    "fields": [stored],
                    "verdict": {
                        "accepted": False,
                        "blocking_reasons": [{"source": "field:vol", "label": "объём партии"}],
                    },
                },
                auto_rejected=True,
                auto_rejection_reason="Авто: объём партии",
                analysis_profile_snapshot=datetime.fromisoformat(saved["updated_at"]),
            )
            return pid
        finally:
            await db.dispose()

    pid = asyncio.run(_seed())
    before = client.get(f"/api/procurements/{pid}").json()
    assert before["status"] == "rejected"
    assert before["auto_rejected"] is True
    assert before["analysis_stale"] is False

    _put("200")
    status = _wait_recheck()
    assert status["total"] >= 1
    assert status["done"] == status["total"]
    assert status["error"] is None

    after = client.get(f"/api/procurements/{pid}").json()
    [field] = after["rag_report"]["fields"]
    assert field["value"] == 300.0
    assert field["condition"]["value"] == "200"
    assert field["match"] is True
    assert after["rag_report"]["verdict"]["accepted"] is True
    assert after["status"] == "new"
    assert after["auto_rejected"] is False
    assert after["analysis_stale"] is False

    # Переименование поля — нужно повторное извлечение: отчёт устаревает.
    renamed = {**_field("200"), "name": "объём вывоза"}
    resp = client.put(
        f"/api/clients/{profile_id}",
        json={"name": active["name"], "competencies": COMP_JSON, "report_fields": [renamed]},
    )
    assert resp.status_code == 200, resp.text
    assert _wait_recheck()["stale"] >= 1
    assert client.get(f"/api/procurements/{pid}").json()["analysis_stale"] is True


def test_site_source_crawl_status_and_text(api_client: tuple[TestClient, Path]) -> None:
    """Сайт-источник: POST ставит сбор, GET показывает ход и итог, текст ищется
    так же, как значения условий (код — в любой записи)."""
    client, _ = api_client
    manager = client.app.state.parser.source_crawls  # type: ignore[attr-defined]

    class _TwoPages:
        pages = ["Шапка\nкод 1 11 010 21 49 2\nПодвал", "Шапка\nкод 4 71 101 01 52 1\nПодвал"]

        def __init__(self) -> None:
            self.i = 0

        async def __aenter__(self) -> _TwoPages:
            return self

        async def __aexit__(self, *exc: object) -> None:
            return None

        async def open(self, url: str) -> None:
            return None

        async def text(self) -> str:
            return self.pages[self.i]

        async def fingerprint(self) -> str:
            return self.pages[self.i]

        async def current_url(self) -> str:
            return f"https://example.org/p{self.i + 1}"

        async def find_next(self, page_no: int) -> Any:
            from zakupki_parser.sources.crawler import NextStep

            return NextStep(mode="number") if self.i == 0 else None

        async def click_next(self) -> None:
            self.i += 1

        async def scroll_to_bottom(self) -> None:
            return None

        async def wait_change(self, old: str, timeout_s: float) -> str | None:
            return self.pages[self.i] if self.pages[self.i] != old else None

    original = manager._driver_factory  # noqa: SLF001
    manager._driver_factory = _TwoPages  # noqa: SLF001
    try:
        created = client.post("/api/sources", json={"url": "https://Example.org/list/"})
        assert created.status_code == 200, created.text
        source_id = created.json()["id"]

        status: dict[str, Any] = {}
        for _ in range(100):
            status = client.get(f"/api/sources/{source_id}").json()
            if not status["active"] and status["status"] != "pending":
                break
            time.sleep(0.05)
        assert status["status"] == "complete"
        assert status["stop_reason"] == "no_next"
        assert status["pages"] == 2
        assert status["text_complete"] is True
        assert status["progress"]["pages"] == 2

        # Тот же сайт в другой записи URL — та же запись, сбор не повторяется.
        again = client.post("/api/sources", json={"url": "https://example.org/list"})
        assert again.json()["id"] == source_id
        assert again.json()["active"] is False

        found = client.get(f"/api/sources/{source_id}/text", params={"q": "47110101521"})
        assert found.status_code == 200
        assert found.json()["matches"] == 1
        assert "4 71 101 01 52 1" in found.json()["fragments"][0]
        head = client.get(f"/api/sources/{source_id}/text").json()
        assert "=== page 1: https://example.org/p1 ===" in head["fragments"][0]
    finally:
        manager._driver_factory = original  # noqa: SLF001


def test_site_source_unsafe_url_rejected(api_client: tuple[TestClient, Path]) -> None:
    from zakupki_parser.net_safety import UnsafeUrlError

    client, _ = api_client
    manager = client.app.state.parser.source_crawls  # type: ignore[attr-defined]

    async def reject(url: str) -> None:
        raise UnsafeUrlError("URL указывает на внутренний/локальный адрес")

    original = manager._check_url  # noqa: SLF001
    manager._check_url = reject  # noqa: SLF001
    try:
        resp = client.post("/api/sources", json={"url": "http://127.0.0.1/admin"})
        assert resp.status_code == 400
        assert "внутренний" in resp.json()["detail"]
    finally:
        manager._check_url = original  # noqa: SLF001


def test_site_source_not_found(api_client: tuple[TestClient, Path]) -> None:
    client, _ = api_client
    assert client.get("/api/sources/999999").status_code == 404


def test_url_condition_rechecked_when_site_collected(api_client: tuple[TestClient, Path]) -> None:
    """Коды ФККО из отчёта против сайта: когда сбор сайта заканчивается, условие
    пересчитывается без LLM; виды работ — из соседнего поля, у каждого кода свои."""
    from scoring_common.conditions import extraction_key, normalize_report_fields
    from zakupki_parser.sources.crawler import NextStep

    client, _ = api_client
    manager = client.app.state.parser.source_crawls  # type: ignore[attr-defined]
    site = {
        "rows": (
            "4 71 101 01 52 1\nлампы\nТранспортирование (1)  Утилизация (1)\nI класс\n"
            "1 11 010 21 49 2\nсемена\nСбор (1)\nII класс"
        )
    }

    class _Site:
        async def __aenter__(self) -> _Site:
            return self

        async def __aexit__(self, *exc: object) -> None:
            return None

        async def open(self, url: str) -> None:
            return None

        async def text(self) -> str:
            return site["rows"]

        async def fingerprint(self) -> str:
            return site["rows"]

        async def current_url(self) -> str:
            return "https://fkko.example.org/org"

        async def find_next(self, page_no: int) -> NextStep | None:
            return None

        async def click_next(self) -> None:
            return None

        async def scroll_to_bottom(self) -> None:
            return None

        async def wait_change(self, old: str, timeout_s: float) -> str | None:
            return None

    active = client.get("/api/clients/active").json()
    profile_id = active["id"]
    fields = [
        {"id": "works", "name": "виды работ", "type": "list"},
        {
            "id": "codes",
            "name": "коды ФККО",
            "type": "list",
            "condition": {
                "op": "all_in",
                "value_kind": "url",
                "value": "https://fkko.example.org/org",
                "near": {"source": "field", "field_id": "works"},
            },
            "blocking": True,
        },
    ]

    def _wait(url: str) -> None:
        for _ in range(200):
            recheck = client.get(f"/api/clients/{profile_id}/recheck").json()
            source = client.post("/api/sources", json={"url": url}).json()
            if not recheck["running"] and not source["active"] and source["status"] != "pending":
                return
            time.sleep(0.05)
        raise AssertionError("сбор/пересчёт не завершился")

    original = manager._driver_factory  # noqa: SLF001
    manager._driver_factory = _Site  # noqa: SLF001
    try:
        saved = client.put(
            f"/api/clients/{profile_id}",
            json={"name": active["name"], "competencies": COMP_JSON, "report_fields": fields},
        )
        assert saved.status_code == 200, saved.text
        _wait("https://fkko.example.org/org")

        # Анализ получает сведения о сайте вместе с условием.
        internal = client.get(
            "/api/clients/active", headers={**INTERNAL_HEADERS, "X-Profile-ID": str(profile_id)}
        ).json()
        meta = next(f for f in internal["report_fields"] if f["id"] == "codes")["condition"][
            "source"
        ]
        assert meta["status"] == "complete" and meta["text_complete"] is True

        # Отчёт закупки: ТЗ требует утилизацию для семян — на сайте её у семян нет.
        defs = {d["id"]: d for d in normalize_report_fields(fields)}
        report_values = [
            {
                "field_id": "works",
                "field_name": "виды работ",
                "field_type": "list",
                "found": True,
                "value": ["утилизация"],
                "extraction_key": extraction_key(defs["works"]),
            },
            {
                "field_id": "codes",
                "field_name": "коды ФККО",
                "field_type": "list",
                "found": True,
                "value": ["4 71 101 01 52 1", "1 11 010 21 49 2"],
                "extraction_key": extraction_key(defs["codes"]),
                "tz_windows": {},
            },
        ]

        async def _seed() -> int:
            db = Database(DbConfig(dsn=TEST_DSN, enabled=True))
            await db.connect()
            try:
                repo = ProcurementRepository(db)
                assert await repo.upsert(
                    {"number": "URL-COND-1", "platform_id": "zakupki_mos", "subject": "Отходы"}
                )
                rows, _ = await repo.list_procurements(number="URL-COND-1")
                pid = rows[0].id
                await repo.upsert_score(
                    pid,
                    profile_id,
                    score_method="fit",
                    rag_report={"status": "ok", "fields": report_values},
                    analysis_profile_snapshot=datetime.fromisoformat(saved.json()["updated_at"]),
                )
                return pid
            finally:
                await db.dispose()

        pid = asyncio.run(_seed())
        source_id = client.post(
            "/api/sources", json={"url": "https://fkko.example.org/org"}
        ).json()["id"]
        # Пересбор сайта -> по окончании пересчёт условий профиля.
        assert client.post(f"/api/sources/{source_id}/refresh").status_code == 200
        _wait("https://fkko.example.org/org")
        card = client.get(f"/api/procurements/{pid}").json()
        codes = next(f for f in card["rag_report"]["fields"] if f["field_id"] == "codes")
        assert codes["match"] is False
        assert codes["mismatch_reasons"] == {"1 11 010 21 49 2": "нет рядом: утилизация"}
        assert card["auto_rejected"] is True
        assert card["analysis_stale"] is False

        # Сайт обновился: у семян появилась утилизация — отклонение снимается.
        site["rows"] = site["rows"].replace("Сбор (1)\nII", "Сбор (1)  Утилизация (1)\nII")
        assert client.post(f"/api/sources/{source_id}/refresh").status_code == 200
        _wait("https://fkko.example.org/org")
        card = client.get(f"/api/procurements/{pid}").json()
        codes = next(f for f in card["rag_report"]["fields"] if f["field_id"] == "codes")
        assert codes["match"] is True
        assert card["auto_rejected"] is False
    finally:
        manager._driver_factory = original  # noqa: SLF001
